import os,json,hashlib,platform
from pathlib import Path
import numpy as np,pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent;RES=ROOT/'results';FIG=ROOT/'figures';FIG.mkdir(exist_ok=True)
navy='#1B3A5C'; blue='#446E9A'; red='#B24A48'; grey='#687789'
D=pd.read_csv(RES/'regrouping_mean_comparison.csv'); V=pd.read_csv(RES/'harmonic_cross_bearing_similarity.csv'); C=pd.read_csv(RES/'frozen_bearing_recall.csv'); O=pd.read_csv(RES/'outer_BPFO_contrast_shift.csv')
print('GROUPS\n', D.to_string(index=False))
fig,ax=plt.subplots(figsize=(9.2,3.4));xs=np.arange(len(D));w=.34
ax.bar(xs-w/2,D.snhop_pct,w,label='SN-HOP',color=blue);ax.bar(xs+w/2,D.full_order_pct,w,label='Full order',color=red)
ax.set_xticks(xs,[('Original' if z=='canonical' else 'Regroup '+str(i)) for i,z in enumerate(D.scheme)]);ax.set_ylabel('Balanced accuracy (%)');ax.set_ylim(0,80);ax.grid(axis='y',alpha=.2);ax.set_axisbelow(True);ax.legend(frameon=False,ncol=2,loc='upper right')
for p in ax.patches: ax.annotate(f'{p.get_height():.1f}',(p.get_x()+p.get_width()/2,p.get_height()+1),ha='center',fontsize=8,color=navy)
fig.tight_layout();fig.savefig(FIG/'FigS2_1_grouping.png',dpi=230);fig.savefig(FIG/'FigS2_1_grouping.svg');plt.close(fig)
fig,ax=plt.subplots(figsize=(9.2,3.4));names=['full_order_256','physical_harmonic_context','uniform_harmonic_context','physical_harmonic_only','broad_context_only'];xlabels=['Full order','Physics + context','Uniform + context','Harmonics only','Broad context']
z=V.set_index('representation').loc[names]
bars=ax.bar(xlabels,z.similarity_gap.values,color=[grey,blue,red,blue,grey],alpha=.9)
for p in bars:ax.annotate(f'{p.get_height():.3f}',(p.get_x()+p.get_width()/2,p.get_height()+.003),ha='center',fontsize=8)
ax.set_ylabel('Cross-bearing class-similarity gap (cosine)');ax.set_ylim(0,.21);ax.grid(axis='y',alpha=.2);ax.set_axisbelow(True);fig.tight_layout();fig.savefig(FIG/'FigS2_2_similarity.png',dpi=230);fig.savefig(FIG/'FigS2_2_similarity.svg');plt.close(fig)
fig,ax=plt.subplots(figsize=(8.5,3.6));names=['KA04','KA15','KA16','KA22','KA30'];x=np.arange(5)
a=C[(C.bearing.isin(names))&(C.direction=='1500_to_900')].set_index('bearing').loc[names].recall.to_numpy()*100
b=C[(C.bearing.isin(names))&(C.direction=='900_to_1500')].set_index('bearing').loc[names].recall.to_numpy()*100
ax.bar(x-.18,a,.36,color=blue,label='1500 to 900 rpm');ax.bar(x+.18,b,.36,color=red,label='900 to 1500 rpm')
ax.set_xticks(x,names);ax.set_ylim(0,110);ax.set_ylabel('Outer-race recall (%)');ax.grid(axis='y',alpha=.2);ax.set_axisbelow(True);ax.legend(frameon=False,ncol=2);fig.tight_layout();fig.savefig(FIG/'FigS2_3_outer_failure.png',dpi=230);fig.savefig(FIG/'FigS2_3_outer_failure.svg');plt.close(fig)

font='Microsoft YaHei'
doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.66);sec.bottom_margin=Inches(.64);sec.left_margin=Inches(.78);sec.right_margin=Inches(.78)
styles=doc.styles;sty=styles['Normal'];sty.font.name=font;sty._element.rPr.rFonts.set(qn('w:eastAsia'),font);sty.font.size=Pt(9.5);sty.paragraph_format.space_after=Pt(3)
for h,size in [('Title',17),('Heading 1',12.5),('Heading 2',10.5)]:
    z=styles[h];z.font.name=font;z._element.rPr.rFonts.set(qn('w:eastAsia'),font);z.font.size=Pt(size);z.font.color.rgb=RGBColor(27,58,92)

def para(text='',bold=False):
    p=doc.add_paragraph();r=p.add_run(text);r.bold=bold;return p

def heading(text,level=1):doc.add_heading(text,level)

def table(columns,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(columns));t.style='Light Shading Accent 1';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for cell,s in zip(t.rows[0].cells,columns):cell.text=str(s)
    for row in rows:
        cells=t.add_row().cells
        for cell,s in zip(cells,row):cell.text=str(s)
    for row in t.rows:
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in c.paragraphs:
                for r in p.runs:r.font.size=Pt(8.1)
    return t

def picture(path,caption):
    doc.add_picture(str(FIG/path),width=Inches(4.95));p=doc.add_paragraph(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.style='Caption'

doc.add_heading('SN-HOP论文返修：第二阶段补充分析报告',0)
para('面向 IEEE Access（Access-2026-47316）的数据复核与补充实验 | 2026年10月19日')
para('证据状态：基于已保存的原始特征缓存、逐条预测与冻结的重建版canonical实现；本报告为返修研究资料，不是投稿终稿。')
heading('一、工作范围与边界')
para('本阶段针对审稿人1的第4、5、7、8条及审稿人2的第1、5条意见，新增：（1）物理轴承重新分组敏感性验证；（2）谐波特征跨轴承、跨转速相似度对照；（3）外圈故障分轴承误分类归因线索；（4）阶次侧带的初步描述性分析。所有分类结果均来自实际执行并保存的逐条预测，不追认历史未保存实验。')
para('必须强调：15个独立物理轴承仍是当前研究的主要样本规模；20项任务共享轴承和训练集，不能作为20个独立物理实验。任何分组抽样、录音窗口或轴承对的重复计数均不可替代新的物理试样。')
heading('二、替代物理轴承分组敏感性')
para('使用原始5组三分类平衡轴承组合，以及4组预先以固定随机种子产生的替代分组（20261020—20261023）。每组保持5个健康、5个外圈故障、5个内圈故障实体不变，只重新安排跨类别的triplet对应关系。每个方案仍采用10种两组留出及双向跨转速，即20项测试；每项的特征设置和分类器均通过其源域3折轴承组验证选择。完整阶次谱基线使用与原始canonical相同的候选规范和嵌套策略。')
rows=[]
for i,r in D.iterrows():rows.append([('原始组合' if i==0 else '替代组合 '+str(i)),f'{r.snhop_pct:.2f}%',f'{r.full_order_pct:.2f}%',f'{r.gain_pp:+.2f} pp'])
table(['组合','SN-HOP BA','完整阶次谱 BA','差值'],rows)
picture('FigS2_1_grouping.png','图S2-1  不同实体轴承分组下的平均平衡准确率')
para(f'四组替代方案中，SN-HOP的BA范围为{D[D.scheme!="canonical"].snhop_pct.min():.2f}%—{D[D.scheme!="canonical"].snhop_pct.max():.2f}%，相对完整阶次谱的差值范围为{D[D.scheme!="canonical"].gain_pp.min():+.2f}至{D[D.scheme!="canonical"].gain_pp.max():+.2f}个百分点。' + ('所有已测分组相对优势均为正，但这一结论仅适用于这5种选定分组和15个轴承，并非跨数据集性能保证。' if (D.gain_pp>0).all() else '部分分组未保持正增益，因此结论必须进一步弱化。'))
heading('三、不同物理轴承的谐波表征相似度')
para('预先固定谐波池化宽度Δ=0.4阶、最大谐波位置35阶（源论文已规定的候选网格值，不依据本次目标结果调参）。针对每个轴承×转速汇总20条录音的特征中心，比较“900 rpm轴承”和“1500 rpm不同轴承”的跨转速余弦相似度，分别统计同一故障类别与不同故障类别，并以两组均值之差作为描述性分离指标。只对展示特征进行全局无标签均值中心化，该处理不进入任何分类器。')
rows=[]
for r in V.itertuples(index=False):rows.append([r.representation,str(r.dimensions),f'{r.same_fault_cross_bearing_cross_speed_cosine_mean:.3f}',f'{r.different_fault_cross_bearing_cross_speed_cosine_mean:.3f}',f'{r.similarity_gap:.3f}'])
table(['特征','维数','同类均值','异类均值','分离度'],rows)
picture('FigS2_2_similarity.png','图S2-2  固定特征设置下的跨轴承跨转速类间相似度差')
para('在固定的描述性比较条件下，“物理谐波+上下文”组的相似度分离度高于完整阶次谱以及相同维数的均匀中心组。其方向支持物理谐波位置具有信息价值，但故障标签参与统计分组，轴承对之间存在依赖关系；因此本结果是直接表征证据，而非独立统计检验或性能优势的因果证明。')
heading('四、外圈故障误分类具体表现')
para('从冻结的SN-HOP基线预测中发现，900→1500 rpm方向下KA15、KA22和KA30的外圈召回率均为0，但错分方向不同。')
CC=pd.read_csv(RES/'frozen_bearing_confusions.csv');rows=[]
for b in ['KA04','KA15','KA16','KA22','KA30']:
    q=CC[(CC.bearing==b)&(CC.direction=='900_to_1500')]; cnt=q.set_index('y_pred').n.to_dict();rows.append([b,str(cnt.get(0,0)),str(cnt.get(1,0)),str(cnt.get(2,0))])
table(['实际外圈轴承','预测健康','预测外圈','预测内圈'],rows)
picture('FigS2_3_outer_failure.png','图S2-3  五个外圈损伤轴承在双向迁移下的召回率差异')
para('每个轴承有20条独立录音，但在4项共享源实体的测试组合中重复预测，因此表中80个分类次数不能视作80个独立物理样本。KA15/KA30主要误判为内圈故障；KA22同时误判为健康和内圈故障。')
doc.add_page_break();heading('BPFO一阶局部对比度核查',2)
A=O[(O.bearing.isin(['KA04','KA15','KA16','KA22','KA30']))&(O.harmonic==1)].copy();rows=[]
for r in A.itertuples(index=False): rows.append([r.bearing,f'{getattr(r,"_2",float("nan")):.3f}' if False else f'{r[2]:.3f}',f'{r[3]:.3f}',f'{r[4]:+.3f}'])
table(['轴承','900 rpm BPFO一阶对比度','1500 rpm BPFO一阶对比度','高低转速差'],rows)
para('一阶BPFO局部对比度在KA15、KA22、KA30较弱，其中KA15与KA22在1500 rpm呈负局部对比度；但KA30一阶对比度变化很小却完全漏检，说明不能把分类失败归因为单一峰强变化。损伤形貌及感应路径解释需要具体损伤元信息或更深入的受控实验支持。')
heading('五、阶次侧带：已作观察，暂不作为提升准确率的证据')
para('以BPFI/BPFO的1—3次谐波为中心，提取±1阶侧点（对应轴转频率偏移）并计算局部峰与中心峰的描述性差值。原始阶次网格步长约0.196阶，且各故障峰、侧点及有限分辨率窗口之间可能发生邻近/重叠，因此此计算仅提供审稿人要求讨论侧带的线索；尚不足以确认真实调制机制，也未对加入侧带特征后是否提升性能进行验证。详见数据文件。')
heading('六、对稿件和回复信的明确修改指向')
for item in [
 'R1-4：在实验设置与局限性中增写多种实体轴承分组的稳定性结果，并承认实体样本只有15个。',
 'R1-5：删除将20个共享轴承测试任务理解为独立样本的强显著性表述；仅将任务差异视为配对描述和内部敏感性证据。',
 'R1-7：新增跨轴承相似度与BPFO谱峰/局部对比度的直接分析，注明它们是表征关联而不是特征抑制机理的因果验证；讨论侧带遗漏。',
 'R1-8 / R2-5：将KA15、KA22和KA30的准确误判去向放入错误分析，基于实体轴承而不是只展示平均混淆矩阵。',
 'R2-1：明确替代分组无法补足外部数据集验证，更广泛轴承几何与工况迁移仍未经验证。',
]: para('• '+item)
heading('七、审计与下一阶段')
para('源码、固定分组配置、所有新增逐项预测、源域参数选择、绘图数据、描述性特征表、运行日志均在第二阶段ZIP。原始600条MAT及原始canonical v1未被改动或覆盖。')
para('下一步：补做非白噪声实验，更新英文正文，完成16条逐点回复、投稿干净稿及黄色高亮PDF；同时如实报告单目标轴承CORAL适配失败的结果。')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('SN-HOP · IEEE Access · Revision evidence · Stage 2')
P=ROOT/'SNHOP_IEEE_Access_返修阶段2实验报告.docx';doc.save(P);print('REPORT',P)
