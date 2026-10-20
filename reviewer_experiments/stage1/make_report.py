import json,datetime,hashlib
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent;O=R/'results'
N=RGBColor(27,58,92);RRED=RGBColor(160,49,54);GREY=RGBColor(84,96,109)
cluster=pd.read_csv(O/'bearing_cluster_sensitivity.csv')
matched=pd.read_csv(O/'matched_dimension_tasks.csv').groupby('rep').target_ba.mean()*100
uda=pd.read_csv(O/'single_target_uda_tasks.csv').groupby('variant').target_ba.mean()*100
window=pd.read_csv(O/'window_resolution_tasks.csv').groupby('window').target_ba.mean()*100
background=pd.read_csv(O/'background_width_tasks.csv').groupby('background_halfwidth').target_ba.mean()*100
band=pd.read_csv(O/'bandpass_tasks.csv').groupby('band').target_ba.mean()*100
timing=json.loads((O/'end_to_end_timing_summary.json').read_text())

d=Document(); sec=d.sections[0];sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.68);sec.left_margin=Inches(.83);sec.right_margin=Inches(.83)
styles=d.styles
normal=styles['Normal'];normal.font.name='Noto Sans CJK SC';normal.font.size=Pt(9); normal.font.color.rgb=RGBColor(36,50,64);normal.paragraph_format.space_after=Pt(5)
for name,sz in [('Title',17),('Heading 1',12),('Heading 2',10)]:
 s=styles[name];s.font.name='Noto Sans CJK SC';s.font.size=Pt(sz);s.font.bold=True;s.font.color.rgb=N;s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(5)

def para(t='',bold=False,color=None):
 p=d.add_paragraph();rr=p.add_run(t);rr.bold=bold
 if color:rr.font.color.rgb=color
 return p

def header(t):d.add_heading(t,level=1)
def sub(t):d.add_heading(t,level=2)
def table(headers,rows,widths=None):
 tbl=d.add_table(rows=1,cols=len(headers));tbl.style='Table Grid';tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,h in enumerate(headers):
  c=tbl.rows[0].cells[i];c.text=str(h)
  tcPr=c._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'1B3A5C');tcPr.append(shd)
  for r in c.paragraphs[0].runs:r.font.color.rgb=RGBColor(255,255,255);r.bold=True;r.font.size=Pt(8)
 for n,row in enumerate(rows):
  ccs=tbl.add_row().cells
  for i,v in enumerate(row):
   c=ccs[i];c.text=str(v);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(1)
    for rr in p.runs:rr.font.size=Pt(8)
   if n%2==1:
    tcPr=c._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'F1F5F9');tcPr.append(shd)
 tbl.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 d.add_paragraph().paragraph_format.space_after=Pt(0)
 return tbl

p=d.add_paragraph(style='Title');p.add_run('SN-HOP · IEEE Access 返修实验阶段报告')
p=d.add_paragraph();p.add_run('2026-10-19  |  Access-2026-47316  |  阶段性补充证据（非最终投稿稿）').font.color.rgb=GREY
para('原始论文：Physics-Guided Speed-Normalized Harmonic Order Pooling for Bearing Fault Diagnosis Across Speeds and Unseen Bearing Specimens')

header('1. 数据与可复现性核验')
para('当前恢复并实际校验了全部15个Paderborn原始轴承ZIP。15/15文件大小与SHA-256匹配，600/600个论文所需MAT记录齐全。新分析均单独保存于stage1_v2目录，未覆盖canonical v1的缓存、参数和预测。')
para('实验基础：15个独立物理轴承；每种速度下300条4秒记录；900/1500 rpm两个方向；10组留出组合共20项迁移任务。原方法严格DG平均平衡准确率为56.67%。')

header('2. 新增对照：维数与谐波定位（R1-6）')
table(['方法','平均平衡准确率','比较约束'],[['SN-HOP（原）','56.67%','源轴承嵌套选择，26–64维'],['完整阶次谱 + PCA',f'{matched["PCA"]:.2f}%','每项任务采用与SN-HOP相同维数'],['完整阶次谱 + Fisher',f'{matched["Fisher"]:.2f}%','每项任务采用与SN-HOP相同维数']])
para('PCA与Fisher降维及分类器选择均仅使用源轴承；特征维数沿用SN-HOP在该任务内通过源域验证确定的维数。当前结果支持“收益不能完全用维数差异解释”，但不代表已完成与完整PFN或全部先进DG方法的公平比较。')

header('3. 参数敏感性（R1-2、R1-3）')
sub('3.1 窗口长度与分辨率')
table(['FFT窗口','900rpm阶次分辨率','1500rpm阶次分辨率','平均平衡准确率'],[['8192（原）','0.521','0.313','56.67%'],['16384','0.260','0.156',f'{window[16384]:.2f}%'],['32768','0.130','0.078',f'{window[32768]:.2f}%']])
para('三种设置均覆盖原始信号前253,952个采样点，帧移固定8192；较长窗口因此产生重叠。每一窗口设置均重新在源轴承内部选择池化参数与分类器。观察到16384性能基本持平，32768下降，不能据此声称原始8192的物理频谱分辨率已经充分。')
sub('3.2 背景阶次半宽')
table(['背景半宽','0.8阶','1.2阶（原）','1.6阶','2.0阶'],[['平均平衡准确率']+[f'{background[v]:.2f}%' for v in (.8,1.2,1.6,2.0)]])
sub('3.3 带通滤波范围')
table(['预设带通范围','200–12,000 Hz','500–10,000 Hz（原）','1,000–8,000 Hz'],[['平均平衡准确率',f'{band["200_12000"]:.2f}%','56.67%',f'{band["1000_8000"]:.2f}%']])
para('两种新增带通配置全部从原始MAT重新预处理600条记录，并在源轴承内重做模型选择。测试范围有限，不能说明500–10,000 Hz在所有设备和工况下最优。')

header('4. 不独立任务的统计敏感性（R1-5）')
table(['对照','均值差（百分点）','5组Bootstrap 95%区间','双侧符号翻转p'],[[r.baseline,f'{r.mean_difference_pp:+.2f}',f'[{r.triplet_boot_CI_low_pp:.2f}, {r.triplet_boot_CI_high_pp:.2f}]',f'{r.exact_5cluster_sign_flip_p_two_sided:.4f}'] for r in cluster.itertuples()])
para('分析先在15个物理轴承上汇总跨测试配置的召回差，再聚合到五个预设的“健康/外圈/内圈”平衡组。只有5个组，且训练集合仍有重叠，因此区间和符号翻转检验均应视为敏感性证据，而非严格独立样本显著性证明。旧稿按20项任务给出的Wilcoxon p值不宜作为主要显著性结论。',color=RRED)

header('5. 单目标轴承UDA与失败分析（R1-8、R1-9）')
table(['场景（每目标轴承10条无标签适配记录）','平均平衡准确率'],[['NoDA（不适配）',f'{uda["NoDA"]:.2f}%'],['逐目标轴承CORAL',f'{uda["CORAL_single_target"]:.2f}%'],['逐目标轴承标准化',f'{uda["TargetZ_single_target"]:.2f}%']])
para('测试使用每个目标轴承的第11–20次测量。CORAL针对单个轴承的10条无标签数据估计目标协方差后重新拟合源分类器。结果明显低于NoDA，说明多目标轴承UDA结果不可直接推广到单目标轴承部署。NoDA 57.75%对应半数保留测试记录，与原56.67%完整测试集不能直接作增益比较。')
para('从原有逐条预测复算：900→1500 rpm方向上，KA15、KA22、KA30三个外圈故障轴承在各自4组留出组合的80条测试出现中，正确识别率均为0%。应将其作为失败机制分析重点，不能预先认定损伤形貌是唯一原因。')

header('6. 完整流水线计时（R1-11、R2-4）')
t=timing['timing_ms'];table(['步骤','中位耗时','95百分位耗时'],[[name,f'{t[k]["median"]:.2f} ms',f'{t[k]["p95"]:.2f} ms'] for name,k in [('读取ZIP及解析MAT','read_decode_ms'),('滤波/Hilbert/FFT/阶次化','filter_envelope_fft_ms'),('谐波池化','pooling_ms'),('分类器预测','classifier_ms'),('合计','full_inference_ms')]])
para('对全部600条原始4秒记录测量；示例分类器在源轴承训练后固定。该基准是在本次云端单线程CPU、已本地缓存ZIP条件下得到，不包括传感器采集、网络传输或模型训练，不代表嵌入式工况的实际延迟。600条重新预处理得到的阶次缓存与v1逐条完全一致。')

header('7. 下一阶段尚需完成')
para('（1）R1-4：固定事先声明的替代轴承分组，完整重跑严格DG验证分组稳定性。\n（2）R1-7：跨轴承谐波特征可视化、类内/类间差异与边带作用。\n（3）R1-8：定位外圈故障错误类别，结合可核查的轴承损伤记录和信号分析。\n（4）R1-10/R1-11：酌情扩展工况或真实工业扰动，无法做到的范围在局限中明确。\n（5）统一更新主文与16条审稿意见回复，形成clean DOCX/PDF、highlighted PDF及最终归档ZIP。')
para('注意：本阶段包为补充实验证据，不是完成返修的投稿稿；v1原始源码属于历史代码丢失后重建的canonical版本，新增结果基于其固定实现。')

foot=sec.footer.paragraphs[0];foot.text='SN-HOP  |  Revision evidence · Stage 1  |  Generated 2026-10-19';foot.style='Caption'
pth=R/'SNHOP_IEEE_Access_返修阶段1实验报告.docx';d.save(pth);print(pth)
