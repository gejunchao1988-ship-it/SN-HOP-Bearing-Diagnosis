from pathlib import Path
import pandas as pd,numpy as np, json, re, hashlib, zipfile, textwrap
import matplotlib;matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent;RES=ROOT/'results'
summary=pd.read_csv(RES/'nonwhite_noise_summary.csv')
pred=pd.read_csv(RES/'nonwhite_noise_predictions.csv')
NAVY=RGBColor(27,58,92);RED=RGBColor(168,47,52);GREY=RGBColor(85,97,113)


def style(doc):
 sec=doc.sections[0];sec.top_margin=Inches(.73);sec.bottom_margin=Inches(.72);sec.left_margin=Inches(.82);sec.right_margin=Inches(.82)
 normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(9);normal.paragraph_format.space_after=Pt(4)
 for n,fs in [('Title',17),('Heading 1',12),('Heading 2',10.4)]:
  s=doc.styles[n];s.font.name='Arial';s.font.size=Pt(fs);s.font.bold=True;s.font.color.rgb=NAVY;s.paragraph_format.space_before=Pt(9);s.paragraph_format.space_after=Pt(5)
 foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
 foot.add_run('SN-HOP IEEE Access Revision  |  Research draft · 19 October 2026').font.size=Pt(8)

def put_table(doc,headers,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for c,v in zip(t.rows[0].cells,headers):
  c.text=str(v);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 for row in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,row):c.text=str(v)
 return t

def make_chart():
 levels=[('Clean',56.6667),('Colored 10 dB',56.25),('Impacts 10 dB',47.1458),('Colored 0 dB',34.625),('Impacts 0 dB',36.9167)]
 fig,ax=plt.subplots(figsize=(7.8,3.15))
 cols=['#1B3A5C','#527293','#527293','#A83B46','#A83B46']
 bars=ax.barh([x[0] for x in levels][::-1],[x[1] for x in levels][::-1],color=cols[::-1],height=.64)
 ax.set_xlim(0,68);ax.set_xlabel('Mean balanced accuracy (%)');ax.axvline(33.333,color='#888888',linestyle=':',linewidth=1,label='Chance (3 classes)')
 for b in bars:ax.text(b.get_width()+.6,b.get_y()+b.get_height()/2,f'{b.get_width():.2f}%',va='center',fontsize=9)
 ax.spines[['top','right']].set_visible(False);ax.legend(frameon=False,loc='lower right',fontsize=8)
 fig.tight_layout();fig.savefig(ROOT/'figures'/'FigS3_1_nonwhite_noise.png',dpi=240,bbox_inches='tight');fig.savefig(ROOT/'figures'/'FigS3_1_nonwhite_noise.svg',bbox_inches='tight');plt.close(fig)


def make_cn():
 d=Document();style(d);d.add_heading('SN-HOP · IEEE Access 返修第三阶段实验报告',0)
 p=d.add_paragraph('实验进度：第三阶段已核验完成（2026-10-19）。本文件是返修证据报告，并非已完成的投稿论文。');p.runs[0].bold=True
 d.add_heading('1. 本阶段任务与实验边界',1)
 d.add_paragraph('依据审稿人1第10、11条及审稿人2第4条意见，在全部15个原始Paderborn轴承ZIP上核验现有数据条件，补做合成非白噪声测试。原始600条振动记录、canonical v1参数及模型选择、源域/目标域划分全部冻结。新增干扰仅进入目标测试信号，分类模型仍在干净源域数据上拟合。')
 d.add_paragraph('干扰类型：① 1200–6500 Hz频带受限有色噪声（Butterworth带通后的随机过程）；② 随机时间出现的阻尼振荡冲击（约2.6–4.6 kHz中心频率，脉冲幅度和衰减随机）。测试信噪比为10 dB与0 dB；每组采用两套固定种子（SeedSequence基于20261019、记录索引、干扰类别和重复编号）。本文仅将其称为“合成扰动”，不等同于真实工业现场噪声。')
 d.add_heading('2. 主要结果',1)
 labels={'clean':'无扰动','band_limited_colored':'带限有色噪声','random_damped_impacts':'随机阻尼冲击'}
 rows=[]
 for r in summary.itertuples():
  condition='无扰动' if r.kind=='clean' else labels[r.kind]+' / '+str(int(r.snr_db))+' dB'
  rows.append((condition,f'{100*r.mean_ba:.2f}',f'{100*r.mean_f1:.2f}',int(r.n_task_appearances)))
 put_table(d,['测试条件','Balanced Accuracy (%)','Macro-F1 (%)','任务×重复'],rows)
 d.add_picture(str(ROOT/'figures'/'FigS3_1_nonwhite_noise.png'),width=Inches(6.7))
 d.add_paragraph('图S3-1：在干净源域训练、目标测试信号受到合成干扰情况下，SN-HOP的平均平衡准确率。虚线为三分类随机基准33.33%；任务配置与物理轴承在不同重复中共享，不能将重复当作新的独立样本。')
 d.add_paragraph('10 dB带限有色噪声下56.25%，接近原干净信号56.67%；10 dB随机阻尼冲击下47.15%，下降9.52个百分点。0 dB下有色噪声和冲击分别为34.63%与36.92%，显示强干扰时诊断性能显著退化。上述结论是描述性的，无独立外部工业噪声验证。')
 d.add_heading('3. 原始数据和预测链核验',1)
 val=json.loads((ROOT/'stage3_verification.json').read_text())
 put_table(d,['核验项目','结果'],[
  ('原始轴承压缩包的SHA-256','15/15通过'),('干净模型预测逐条回放','2,400/2,400一致'),('源数据记录','600条'),('合成扰动所得频谱','4,800组'),('新增任务行','160（20任务×2噪声×2 SNR×2重复）'),('新增预测行','19,200（重复测试出现，不是19,200个独立记录）'),('目标注入信噪比最大误差',f'{val["SNR_max_abs_error_dB"]:.2e} dB'),('结果核验脚本','PASS')])
 d.add_heading('4. 新负载工况可行性',1)
 d.add_paragraph('逐一查看目前恢复的15个原始ZIP，每个ZIP只有N09_M07_F10与N15_M07_F10两类真实工况（每类20个MAT）。帕德博恩大学官方数据中心另外记录了N15_M01_F10（1500 rpm/0.1 Nm/1000 N）与N15_M07_F04（1500 rpm/0.7 Nm/400 N）。现有15个ZIP不包含这两类测量，不能直接声称已完成跨负载实验。')
 d.add_paragraph('官方来源： https://mb.uni-paderborn.de/en/kat/research/bearing-datacenter/operating-conditions')
 d.add_paragraph('后续如下载官方完整档案，需要核验15个目标物理轴承身份、各工况20条MAT、相同传感器通道、速度测量，并用全新预设源域/目标域协议评价。不得用人工变速或插值代替真实工况。')
 d.add_heading('5. 返修表述与仍待解决问题',1)
 for t in [
  '审稿意见R1-11：已增加合成非白噪声与完整流水线耗时（阶段1中位数46.24 ms/4秒记录）；明确两者均不等于真实工业系统测试。',
  '审稿意见R1-10：已查明官方新增负载工况，但尚未取得对应原始文件，因此仍未完成跨负载实测；不能将该项标为“已解决”。',
  '现有15个轴承在5套重组中均复用，不能将20个任务和多组实验解释为独立物理样本。',
  '真正投稿前，仍需判断是否有必要实现完整PFN对照、补充代表性DG方法以及将所有新增数据融入英文正文和高亮稿。']:
  d.add_paragraph(t,style='List Bullet')
 d.add_heading('6. 文件与复现说明',1)
 d.add_paragraph('src/run_nonwhite_noise.py从SHA-256校验通过的原始ZIP重新提取所有4秒测量，生成干扰信号与256维阶次频谱；干净源域分类器采用canonical v1源轴承选定配置。src/verify_stage3.py逐条检查预测、指标、SNR、v1完全重放和缓存形状；figures保存PNG/SVG；results保存全部CSV、压缩频谱及核验记录。完整ZIP未复制约2.12 GiB原始文件。')
 d.save(ROOT/'SNHOP_IEEE_Access_返修阶段3实验报告.docx')

COMMENTS=[
(1,'Novelty versus characteristic-order analysis and physical frequency normalization.'),
(2,'Sampling/window spectral resolution 8192/64 kHz and harmonic pooling width 0.2 order.'),
(3,'Selection and sensitivity for 500–10000 Hz band and background halfwidth.'),
(4,'Only 15 physical specimens; per-bearing variation and alternative balanced groupings.'),
(5,'Repeated-bearing dependence in the 20 paired tasks; confidence intervals and statistical inference.'),
(6,'Coverage- and dimension-matched baselines, partial PFN-core, domain generalization comparisons.'),
(7,'Direct cross-bearing harmonic feature evidence and possible sideband contribution.'),
(8,'Outer-race low recall and loss of recall after CORAL.'),
(9,'One unseen target bearing UDA, CORAL versus standardization and lack of significance.'),
(10,'Only two near-constant speed conditions, fixed loading, one bearing geometry.'),
(11,'End-to-end runtime, realistic noise, publicly accessible code/configuration.'),
]
RESP={
1:('We agree that envelope demodulation and order normalization are established. The revised positioning does not claim either operation is new. The operational distinction lies in fixed theoretical BPFO/BPFI harmonic neighborhoods, their peak-versus-local-background contrasts, coarse-band context, and strict source-bearing-only selection. We added equivalent-dimensional full-order controls and direct fixed-feature cross-bearing analyses.', 'Add a contribution comparison to FCOA/PFN with references; clearly note absence of a learned extractor. The complete published PFN system is not yet reproduced.'),
2:('We agree that interpolation does not recover FFT resolution. Using the same 253,952 raw samples, 16,384-point overlapped windows give 56.88% mean balanced accuracy versus 56.67% at 8,192 points, while 32,768 gives 49.00%. We will interpret this as a window/resolution tradeoff, not evidence that 8,192 is physically resolution-perfect.','Insert task-level sensitivity results and explicitly distinguish order-grid sampling from intrinsic FFT resolution.'),
3:('Bandpass alternatives 200–12,000 Hz and 1,000–8,000 Hz reach 52.46% and 49.29% versus 56.67% at 500–10,000 Hz. Alternative background halfwidths 0.8, 1.6, 2.0 orders yield 55.17%, 51.50%, 50.83%, compared with 56.67% for 1.2. All model configuration choices were reselected within source-bearing folds.','Add prespecified local parameter sensitivity table; refrain from claiming a globally optimal band.'),
4:('Across the original and four fixed balanced regroupings, SN-HOP gives 56.67%, 59.38%, 58.88%, 51.17%, and 54.92% mean balanced accuracy; the full-order representation gives 41.33%, 37.00%, 37.33%, 38.33%, and 38.92%. We also tabulate individual-bearing recall and failure patterns. All grouping results share the same 15 specimens and hence do not increase independent sample size.','Add regrouping design, task-level predictions and bearing-specific class recall; retain small-specimen limitations.'),
5:('The initial 20 Wilcoxon task statistics do not represent 20 independent populations. A sensitivity analysis based on five prespecified bearing triplets gives a two-sided exact sign-flip p=0.0625 for SN-HOP over full order despite the +15.33 percentage point observed mean difference. The corresponding triplet bootstrap interval is approximately +7.29 to +27.17 pp but should be interpreted cautiously with five clusters and overlapping tasks.','Reduce inferential claims, report dependent-task uncertainty, and separate descriptive effect sizes from formal evidence; external sample diversity is still needed.'),
6:('Equal-dimensional source-fitted full-order PCA and Fisher-selection controls reach 45.21% and 43.25% versus SN-HOP 56.67%. This strengthens the evidence that low dimensionality alone does not fully explain the result. We recognize that PFN-core (55.42%) is not a faithful reproduction of the complete published PFN system and must not be presented as such.','Add equal-dimensional comparison and accurately label PFN-core; additional matched frequency-coverage and a genuine DG method remain pending before submission.'),
7:('On 30 bearing-by-speed centroids, fixed-parameter physical harmonic pooling gives 0.1585 same-class minus different-class cross-bearing cosine separation, compared with full-order 0.0869 and uniform matched-dimensional pooling 0.0731. These are descriptive and reuse physical bearings, so no causal or independent-sample significance claim is made. Descriptive sideband observations are included with finite-resolution limitations.','Add paired feature/centroid plots and discuss that sidebands could carry useful information but were not proven to improve generalization.'),
8:('In the reverse 900→1500 rpm condition, three outer-race specimens were never correctly classified across four test appearances of their same 20 recordings: KA15 (4 healthy, 76 inner), KA22 (45 healthy, 35 inner), KA30 (0 healthy, 80 inner). This makes the outer-race weakness concrete. It does not prove that damage morphology caused the failures.','Add per-bearing confusion and contrasting harmonic signatures; avoid unsupported damage-morphology causal attributions.'),
9:('For individual-target-bearing unlabeled adaptation, mean balanced accuracy is 42.50% with CORAL and 39.17% with target standardization, below 57.75% without adaptation on the same reserved test subset. The original multi-target adaptation yielded 62.42% for CORAL and 62.50% for standardization, without a meaningful CORAL advantage.','Reframe CORAL as a limited optional comparator, not a deployment benefit for an individual new bearing.'),
10:('The verified current files contain only N09_M07_F10 and N15_M07_F10. The Paderborn University official catalog additionally lists N15_M01_F10 and N15_M07_F04, but those genuine raw records have not been recovered in our current archives. We cannot claim evidence for varying loads or other bearing geometries.','Pending: retrieve and verify the additional real load-condition recordings and evaluate a preregistered bearing-wise protocol, or explicitly delimit the scope and justify why the required validation cannot be provided.'),
11:('An instrumented complete inference pipeline on 600 recordings, including local ZIP read/decode, bandpass filtering, Hilbert/FFT feature construction, harmonic pooling, and prediction, gave 46.24 ms median per four-second recording on the reference hosted CPU; sensor acquisition and transfer are excluded. New synthetic nonwhite perturbations yielded mean balanced accuracy of 56.25%/34.63% for band-limited colored noise at 10/0 dB and 47.15%/36.92% for random damped impulses at 10/0 dB. The strong-noise failures are discussed openly.','Add runtime methods and nonwhite robustness table/figure; provide an accessible code repository link only after publication/upload is actually completed.'),
}
RESP2={
1:('We agree that one dataset and only 15 physical bearings limit external validity. Four balanced regroupings have been completed, but these do not add specimens or a new external dataset.','Report fixed split sensitivity and narrow generalization claims. Additional dataset validation remains pending.'),
2:('The 56.67% mean balanced accuracy in a demanding three-class unseen-bearing cross-speed setting is not sufficient evidence for safety-critical or autonomous industrial decisions. The revised interpretation emphasizes a diagnostic representation benchmark and explicitly discusses erroneous outer-race classifications.','Reframe practical applicability and state that deployment would require broader specimen coverage, calibration, fail-safe logic, and prospective validation.'),
3:('The 1.25 pp observed improvement over PFN-core is not statistically significant and does not establish superiority over the complete published PFN system. The contribution is the interpretable compact representation and leakage-resistant evaluation, with equal-dimension controls and physics-feature analyses.', 'Remove broad superiority language. Additional complete-PFN validation, if possible, is pending.'),
4:('Full inference, not solely final classifier prediction, has now been measured on all 600 locally stored recordings: median 46.24 ms, mean 46.81 ms, and 95th percentile 56.80 ms per four-second recording. These figures depend on the reference CPU and local cache.','Add a complete latency breakdown and hardware/software/environment notes.'),
5:('We have added identity-specific outer-race errors, including KA15, KA22, and KA30, which show 0% recall in the 900→1500 rpm testing direction. The evidence supports a concrete limitation; it does not isolate a unique physical cause.','Expand error analysis and discuss richer sideband/damage morphology features as hypotheses for future work rather than completed improvements.'),
}

def make_reply():
 d=Document();style(d);d.add_heading('Response to Reviewers — WORKING DRAFT',0)
 x=d.add_paragraph('Manuscript: Physics-Guided Speed-Normalized Harmonic Order Pooling for Bearing Fault Diagnosis Across Speeds and Unseen Bearing Specimens');x.runs[0].bold=True
 d.add_paragraph('Manuscript ID: Access-2026-47316 | Date: 19 October 2026')
 p=d.add_paragraph('NOT READY FOR RESUBMISSION: This draft reports verified supplement experiments, but editorial page/line references, full manuscript integration, matched frequency-coverage/DG comparisons and supplementary load experiments remain unresolved. Do not submit as a final author response.');p.runs[0].font.color.rgb=RED;p.runs[0].bold=True
 d.add_heading('General response',1)
 d.add_paragraph('We thank the Editor and both reviewers for the detailed and constructive assessment. We acknowledge that the physical frequency normalization and envelope analysis steps are established tools. Our revisions focus on separating the contributions of harmonic localization and dimensionality, expanding robust cross-bearing evidence, correcting dependent-task statistical interpretation, and reporting operational limitations without overstating field readiness. Below we distinguish completed verifiable work from evidence still required before the resubmission.')
 d.add_heading('Reviewer 1',1)
 for i,comment in COMMENTS:
  d.add_heading(f'Comment R1-{i}',2)
  p=d.add_paragraph();p.add_run('Reviewer concern (summarized): ').bold=True;p.add_run(comment)
  for label,txt in [('Response — verified evidence: ',RESP[i][0]),('Manuscript action / open item: ',RESP[i][1])]:
   p=d.add_paragraph();p.add_run(label).bold=True;p.add_run(txt)
  d.add_paragraph('Location in revised manuscript: [TO BE INSERTED AFTER MANUSCRIPT REVISION]',style=None)
 d.add_heading('Reviewer 2',1)
 for i,comment in [(1,'Data diversity and single bearing dataset.'),(2,'Practical suitability of 56.67% balanced accuracy.'),(3,'Lack of significant superiority over PFN-core.'),(4,'Missing complete preprocessing runtime.'),(5,'Outer-race misclassifications and improvements.')]:
  d.add_heading(f'Comment R2-{i}',2)
  p=d.add_paragraph();p.add_run('Reviewer concern (summarized): ').bold=True;p.add_run(comment)
  for label,txt in [('Response — verified evidence: ',RESP2[i][0]),('Manuscript action / open item: ',RESP2[i][1])]:
   p=d.add_paragraph();p.add_run(label).bold=True;p.add_run(txt)
  d.add_paragraph('Location in revised manuscript: [TO BE INSERTED AFTER MANUSCRIPT REVISION]')
 d.add_heading('Before final submission: mandatory quality gate',1)
 for t in ['Add page and line ranges for every response; ensure each cited experiment, table and figure is actually in the revised manuscript.', 'Resolve the incomplete full-PFN and representative DG method baseline concerns, or provide transparent justified limitations.', 'Decide whether the publicly available additional Paderborn real-load recordings will be downloaded and verified for a new protocol; do not simulate them.', 'Publish code/configuration with a real URL or DOI before claiming public access.', 'Produce clean IEEE Access DOCX/PDF, highlighted PDF including grammar changes, and upload author response in the required category.']:
  d.add_paragraph(t,style='List Bullet')
 d.save(ROOT/'Response_to_Reviewers_WORKING_DRAFT.docx')

if __name__=='__main__':make_chart();make_cn();make_reply();print('REPORT AND RESPONSE WORKING DRAFT GENERATED')
